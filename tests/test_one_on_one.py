"""The AM's 1:1 document: the page, its data, and the Word download."""
from io import BytesIO

import pytest

pytestmark = pytest.mark.usefixtures("fresh_db")


def _am():
    from sqlalchemy import select
    from db import SessionLocal
    from jobs import run_weekly_scoring
    from db import current_week_start
    from datetime import timedelta
    from models import User
    run_weekly_scoring(current_week_start() - timedelta(weeks=1))
    run_weekly_scoring(current_week_start())
    s = SessionLocal()
    return s, s.scalars(select(User).where(User.role == "am")).first()


def test_doc_rows_are_in_status_order():
    import main
    from models import WEEK_STATUS_LABELS
    s, am = _am()
    d = main.one_on_one_doc(s, am)
    order = d["status_order"]
    ranks = [order.index(r["ws"]["status"]) for r in d["rows"]]
    assert d["rows"] and ranks == sorted(ranks)
    assert sum(d["counts"].values()) == len(d["rows"])
    assert WEEK_STATUS_LABELS["blocked"] == "Blocked" and WEEK_STATUS_LABELS["working"] == "Working"
    assert len(d["wins"]) == len(set(d["wins"]))
    s.close()


def test_docx_is_a_word_file_with_the_accounts_and_notes():
    import docs_export
    import main
    from docx import Document
    s, am = _am()
    d = main.one_on_one_doc(s, am)
    doc = Document(BytesIO(docs_export.one_on_one_docx(d, "Ask about promotion")))
    text = "\n".join(p.text for p in doc.paragraphs)
    cells = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert am.name in text and "Ask about promotion" in text and "This week's accounts" in text
    assert all(r["account"].name in cells for r in d["rows"])
    s.close()
