"""Export the demo data as a workbook with one tab per source system.

The tabs hold exactly the rows the app was seeded with, shaped the way each source system would
deliver them, so a reviewer can see what PriorityOne reads and trace any flagged account back to
its raw rows. Two output tabs show what the tracker made of them.

Run after seeding:  python export_sources.py  [path]   (default: priorityone_source_data.xlsx)
"""
import random
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from db import SessionLocal, current_week_start
from contacts import ROLE_LABELS
from models import (Account, AccountPaymentMethod, CallNote, Contact, PrioritySnapshot, SignalEvent, SourceSync,
                    User, WeeklyMetric)
from signals import PAYMENT_METHOD_LABELS

FONT = "Arial"
NAVY = "001222"
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
BODY = Font(name=FONT, size=10)
INPUT = Font(name=FONT, size=10, color="0000FF")        # hardcoded input the reader may change
LINK = Font(name=FONT, size=10, color="008000")         # formula pulling from another tab
URL = Font(name=FONT, size=10, color="0563C1", underline="single")
TITLE = Font(name=FONT, size=14, bold=True)
BOLD = Font(name=FONT, size=10, bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
DATE = "yyyy-mm-dd"
USD = "$#,##0;($#,##0);-"


def sheet(wb, title, headers, rows, formats=None, wrap=(), widths=None, links=(), note=None):
    """A data tab: navy header, frozen, filterable, with per-column number formats."""
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = HEADER_FONT, HEADER_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(list(r))
    formats = formats or {}
    for col_idx, header in enumerate(headers, start=1):
        letter = get_column_letter(col_idx)
        fmt = formats.get(header)
        for (cell,) in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            cell.font = BODY
            if fmt:
                cell.number_format = fmt
            if header in wrap:
                cell.alignment = WRAP
            if header in links and isinstance(cell.value, str) and cell.value.startswith("http"):
                cell.hyperlink, cell.font = cell.value, URL
        longest = max([len(str(header))] + [len(str(r[col_idx - 1])) for r in rows[:300] if r[col_idx - 1] is not None])
        ws.column_dimensions[letter].width = (widths or {}).get(header, min(max(10, longest + 2), 60))
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    if note:
        ws["A1"].comment = Comment(note, "PriorityOne")
    return ws


def build(path: str) -> None:
    week = current_week_start()
    s = SessionLocal()
    accounts = list(s.scalars(select(Account).order_by(Account.name)))
    acc = {a.id: a for a in accounts}
    users = {u.id: u.name for u in s.scalars(select(User))}
    merchant = {a.id: f"ACC-{a.id:05d}" for a in accounts}
    events = list(s.scalars(select(SignalEvent).order_by(SignalEvent.detected_on, SignalEvent.id)))

    wb = Workbook()
    wb.remove(wb.active)
    readme = wb.create_sheet("README")

    # --- Salesforce -------------------------------------------------------------------------
    sheet(wb, "SF Accounts",
          ["Account Id", "Account Name", "Segment", "Vertical", "Region", "ARR (USD)", "Customer Since",
           "Account Owner", "CSM", "Primary Contact", "Previous Primary Contact", "Contact Changed On",
           "Exec Sponsor Last Contact", "Engaged Contacts (90d)", "Record URL"],
          [[a.salesforce_id, a.name, a.segment, a.vertical, a.region, a.arr_usd, a.customer_since, users[a.owner_id],
            a.csm_name, a.primary_contact_name, a.previous_contact_name, a.contact_changed_on,
            a.exec_sponsor_last_contact, a.contacts_engaged_90d, f"https://adyen.my.salesforce.com/{a.salesforce_id}"]
           for a in accounts],
          formats={"ARR (USD)": USD, "Customer Since": DATE, "Contact Changed On": DATE, "Exec Sponsor Last Contact": DATE},
          links=("Record URL",),
          note="Salesforce Account object. Feeds: contact_departed, exec_sponsor_gap, single_threaded, and importance (segment, ARR, tenure).")

    people = list(s.scalars(select(Contact).order_by(Contact.account_id, Contact.is_primary.desc(), Contact.id)))
    sheet(wb, "SF Contacts",
          ["Contact Id", "Account Id", "Account Name", "Name", "Title", "Role (PriorityOne)", "Email", "Primary Contact",
           "Last Activity", "Left Company On", "Record URL"],
          [[c.salesforce_id, acc[c.account_id].salesforce_id, acc[c.account_id].name, c.name, c.title,
            ROLE_LABELS.get(c.role, c.role), c.email, "Yes" if c.is_primary else "", c.last_activity_on, c.left_on,
            f"https://adyen.my.salesforce.com/{c.salesforce_id}"] for c in people],
          formats={"Last Activity": DATE, "Left Company On": DATE}, links=("Record URL",),
          note="Salesforce Contact object, with a role mapped from the title. Feeds the ranked contact list and "
               "email drafts; Gong attendees (Gong calls tab) add call history. Emails use invented .example domains.")

    ws = sheet(wb, "SF Contracts",
               ["Contract Number", "Account Id", "Account Name", "Status", "End Date", "Days to Renewal"],
               [[f"0000{a.id:04d}", a.salesforce_id, a.name, "Activated", a.renewal_date, None] for a in accounts],
               formats={"End Date": DATE}, widths={"Days to Renewal": 16},
               note="Salesforce Contract object. End Date is the renewal date. Feeds: renewal_approaching (≤60 days medium, ≤30 high).")
    for r in range(2, len(accounts) + 2):
        ws[f"F{r}"] = f"=E{r}-'Source freshness'!$B$3"
        ws[f"F{r}"].font, ws[f"F{r}"].number_format = LINK, "0"

    ws = sheet(wb, "SF QBR Events",
               ["Event Id", "Account Id", "Account Name", "Type", "Activity Date", "Days Since"],
               [[f"00U{a.id:06d}QBR", a.salesforce_id, a.name, "QBR", a.last_qbr_date, None] for a in accounts],
               formats={"Activity Date": DATE}, widths={"Days Since": 12},
               note="Salesforce Event records with Type = QBR (latest per account). Feeds: qbr_overdue (>180 days), no_contact_90d (with Gong calls).")
    for r in range(2, len(accounts) + 2):
        ws[f"F{r}"] = f"='Source freshness'!$B$3-E{r}"
        ws[f"F{r}"].font, ws[f"F{r}"].number_format = LINK, "0"

    # --- Payments data warehouse --------------------------------------------------------------
    metrics = list(s.scalars(select(WeeklyMetric).order_by(WeeklyMetric.account_id, WeeklyMetric.week_start)))
    sheet(wb, "Payments weekly",
          ["Merchant Account", "Account Name", "Week Start", "Processing Volume (USD)", "Auth Rate",
           "Chargeback Rate", "Failed Payouts"],
          [[merchant[m.account_id], acc[m.account_id].name, m.week_start, round(m.processing_volume_usd, 2),
            m.auth_rate, m.chargeback_rate, m.failed_payout_count] for m in metrics],
          formats={"Week Start": DATE, "Processing Volume (USD)": USD, "Auth Rate": "0.0%", "Chargeback Rate": "0.00%"},
          note="Payments data warehouse, merchant_weekly_metrics. One row per account per completed week (26 weeks). "
               "Feeds: volume_drop / volume_growth (seasonally adjusted), auth_rate_drop, chargeback_spike, failed_payout.")

    pms = list(s.scalars(select(AccountPaymentMethod).order_by(AccountPaymentMethod.account_id, AccountPaymentMethod.enabled_on)))
    sheet(wb, "Payment methods",
          ["Merchant Account", "Account Name", "Vertical", "Payment Method", "Enabled On"],
          [[merchant[p.account_id], acc[p.account_id].name, acc[p.account_id].vertical,
            PAYMENT_METHOD_LABELS.get(p.method, p.method), p.enabled_on] for p in pms],
          formats={"Enabled On": DATE},
          note="Payments data warehouse, payment_method_config. Feeds: payment_method_gap (≥60% of vertical peers offer a method the account doesn't).")

    # --- Support desk -------------------------------------------------------------------------
    sev1 = [e for e in events if e.signal_type == "sev1_open"]
    sheet(wb, "Support tickets",
          ["Ticket Id", "Merchant Account", "Account Name", "Priority", "Opened On", "Status", "Closed On", "Ticket URL"],
          [[e.source_ref.rsplit("/", 1)[-1], merchant[e.account_id], acc[e.account_id].name, "Sev-1", e.detected_on,
            "Open" if e.resolved_on is None else "Closed", e.resolved_on, e.source_ref] for e in sev1],
          formats={"Opened On": DATE, "Closed On": DATE}, links=("Ticket URL",),
          note="Support desk tickets with priority Sev-1. Feeds: sev1_open (always high severity).")
    sheet(wb, "Support weekly",
          ["Merchant Account", "Account Name", "Week Start", "Open Tickets", "Open Sev-1"],
          [[merchant[m.account_id], acc[m.account_id].name, m.week_start, m.support_tickets_open, m.support_sev1_open]
           for m in metrics],
          formats={"Week Start": DATE},
          note="Support desk, open ticket counts per account per week.")

    # --- Gong ---------------------------------------------------------------------------------
    notes = list(s.scalars(select(CallNote).order_by(CallNote.call_date, CallNote.id)))
    sheet(wb, "Gong calls",
          ["Call Id", "Account Name", "Call Date", "Attendees", "Call Brief", "Risks (extracted)",
           "Opportunities (extracted)", "Sentiment", "Call URL"],
          [[n.source_ref.rsplit("=", 1)[-1], acc[n.account_id].name, n.call_date, n.attendees, n.summary,
            "; ".join(n.extracted_risks), "; ".join(n.extracted_opportunities), n.sentiment, n.source_ref] for n in notes],
          formats={"Call Date": DATE, "Sentiment": "+0.0;-0.0;0.0"}, wrap=("Call Brief", "Attendees"),
          widths={"Call Brief": 70, "Attendees": 40, "Call Id": 22}, links=("Call URL",),
          note="Gong calls with AI briefs, tracker-extracted risks/opportunities and a sentiment score (-1..1). "
               "Feeds: cross_sell_fit ('Interest in …'), sentiment_decline, no_contact_90d. Gong is deliberately 12 days stale.")

    # --- Slack (shared customer channels) -----------------------------------------------------
    slack_rows = []
    for e in events:
        if e.signal_type == "contact_departed" and e.source_system == "slack":
            a = acc[e.account_id]
            slack_rows.append([f"#shared-{a.name.lower().replace(' ', '-').replace('&', 'and')}", a.name, e.detected_on,
                               f"{a.previous_contact_name} ({a.name})",
                               f"Hi all, today is my last day at {a.name}. {a.primary_contact_name} will be your contact for payments from now on. Thanks for everything!",
                               "Contact change", "contact_departed"])
    r = random.Random(42)
    benign = ["Thanks for the quick fix on the webhook delays yesterday!",
              "Can someone confirm when the September settlement report will be ready?",
              "Heads-up: we launch our autumn promo on Monday, so expect higher volume.",
              "Is there a sandbox we can use to test the new 3DS flow?",
              "Great session today, sharing the slides with the team.",
              "Could you re-send the invoice for August? Finance can't find it.",
              "We'd like to add two more users to the Customer Area.",
              "Quick question on payout timing for the new EUR account."]
    for i, text in enumerate(benign):
        a = r.choice(accounts)
        slack_rows.append([f"#shared-{a.name.lower().replace(' ', '-').replace('&', 'and')}", a.name,
                           week.replace(day=1) if i % 2 else week, f"{a.primary_contact_name} ({a.name})", text,
                           "No signal", ""])
    slack_rows.sort(key=lambda x: x[2])
    sheet(wb, "Slack messages",
          ["Channel", "Account Name", "Posted On", "Author", "Message", "Classified As", "Feeds Signal"],
          slack_rows, formats={"Posted On": DATE}, wrap=("Message",), widths={"Message": 70},
          note="Shared Slack Connect channels with customers. Every message is classified; only contact changes feed the "
               "tracker (contact_departed). Messages classified 'No signal' are illustrative and ignored.")

    # --- NPS survey ---------------------------------------------------------------------------
    nps = [a for a in accounts if a.nps_score is not None]
    nps.sort(key=lambda a: a.nps_date)
    ws = sheet(wb, "NPS responses",
               ["Response Id", "Account Name", "Respondent", "Score (0-10)", "Category", "Submitted On"],
               [[f"NPS-{a.id:05d}-{a.nps_date:%y%m%d}", a.name, a.primary_contact_name, a.nps_score, None, a.nps_date]
                for a in nps],
               formats={"Submitted On": DATE}, widths={"Category": 12},
               note="Relationship NPS survey, synced weekly (Fridays). Feeds: nps_detractor (score 0-6 within 90 days; ≤3 high).")
    for row in range(2, len(nps) + 2):
        ws[f"E{row}"] = f'=IF(D{row}>=9,"Promoter",IF(D{row}>=7,"Passive","Detractor"))'
        ws[f"E{row}"].font = BODY

    # --- Ingestion log and freshness ------------------------------------------------------------
    syncs = list(s.scalars(select(SourceSync).order_by(SourceSync.source_system, SourceSync.synced_on)))
    sheet(wb, "Sync log", ["Source", "Synced On", "Rows Synced"],
          [[x.source_system, x.synced_on, x.rows_synced] for x in syncs], formats={"Synced On": DATE},
          note="Ingestion log: one row per successful sync. The tracker uses it to decide whether a source is stale.")

    fr = wb.create_sheet("Source freshness")
    fr["A1"], fr["A1"].font = "Source freshness", TITLE
    fr["A3"], fr["A3"].font = "Week being scored (Monday)", BOLD
    fr["B3"], fr["B3"].font, fr["B3"].number_format = week, INPUT, DATE
    fr["C3"] = "← input: change it to check freshness as of another week"
    fr["C3"].font = Font(name=FONT, size=9, italic=True, color="646F79")
    fr["A4"] = "A source is stale when its newest sync is more than 7 days before the week being scored."
    fr["A4"].font = Font(name=FONT, size=9, italic=True)
    hdr = ["Source", "Newest Sync", "Days Behind", "Stale?", "Feeds"]
    for i, h in enumerate(hdr, start=1):
        c = fr.cell(row=6, column=i, value=h)
        c.font, c.fill = HEADER_FONT, HEADER_FILL
    feeds = {"payments_db": "volume, auth rate, chargebacks, payouts, payment methods",
             "salesforce": "renewals, QBRs, contacts, exec sponsor, threading",
             "gong": "call notes, cross-sell, sentiment, no-contact",
             "support": "sev-1 tickets", "slack": "contact changes", "survey": "NPS"}
    for i, (src, what) in enumerate(feeds.items(), start=7):
        fr[f"A{i}"], fr[f"A{i}"].font = src, BODY
        fr[f"B{i}"] = f"=_xlfn.MAXIFS('Sync log'!$B:$B,'Sync log'!$A:$A,A{i})"
        fr[f"B{i}"].font, fr[f"B{i}"].number_format = LINK, DATE
        fr[f"C{i}"] = f"=$B$3-B{i}"
        fr[f"C{i}"].font, fr[f"C{i}"].number_format = BODY, "0"
        fr[f"D{i}"] = f'=IF(C{i}>7,"Yes","No")'
        fr[f"D{i}"].font = BODY
        fr[f"E{i}"], fr[f"E{i}"].font = what, BODY
    for col, w in zip("ABCDE", (28, 14, 13, 9, 52)):
        fr.column_dimensions[col].width = w

    # --- Tracker output -------------------------------------------------------------------------
    sheet(wb, "Signals (output)",
          ["Account Name", "Signal", "Severity", "Detected On", "Resolved On", "Resolution", "Source System",
           "Summary", "Source Link"],
          [[acc[e.account_id].name, e.signal_type, e.severity, e.detected_on, e.resolved_on, e.resolution,
            e.source_system, e.summary, e.source_ref] for e in events],
          formats={"Detected On": DATE, "Resolved On": DATE}, wrap=("Summary",), widths={"Summary": 70},
          links=("Source Link",),
          note="OUTPUT, not a source: the signal events the tracker detected from the source tabs, with their history "
               "(superseded = replaced by a more severe event; cleared = the condition went away).")

    snaps = list(s.scalars(select(PrioritySnapshot).where(PrioritySnapshot.week_start == week)
                           .order_by(PrioritySnapshot.user_id, PrioritySnapshot.rank)))
    rows = []
    for p in snaps:
        top = next((x for x in p.reasons_json if x["signal_type"] != "resurfaced"), None)
        rows.append([users[p.user_id], p.rank, acc[p.account_id].name, p.category.replace("_", " ").title(),
                     round(p.score, 1), p.revenue_at_risk, top["sentence"] if top else "",
                     "Yes" if p.suppressed else "No", "Yes" if p.partial else "No"])
    sheet(wb, "Weekly lists (output)",
          ["Account Manager", "Rank", "Account Name", "Category", "Score", "Revenue at Risk (USD)", "Top Reason",
           "Hidden After Feedback", "Data Incomplete"],
          rows, formats={"Revenue at Risk (USD)": USD}, wrap=("Top Reason",), widths={"Top Reason": 70},
          note=f"OUTPUT, not a source: every scored account for the week of {week:%d %b %Y}, ranked within each AM's book. "
               "The app shows the top 7 Act Now / Opportunity accounts that aren't hidden; Maintain goes to the hygiene checklist.")

    # --- README -------------------------------------------------------------------------------
    readme["A1"], readme["A1"].font = "PriorityOne: source data", TITLE
    lines = [
        "One tab per source system, holding the exact rows the demo app was seeded with, shaped the way each system would deliver them.",
        "All accounts, people, figures and links are invented mock data.",
        f"Data as of the week of {week:%d %b %Y}. Regenerate after seeding: python export_sources.py",
        "Blue text = an input you can change. Green text = a formula pulling from another tab. Hover a tab's A1 cell for notes.",
    ]
    for i, t in enumerate(lines, start=2):
        readme[f"A{i}"], readme[f"A{i}"].font = t, BODY
    hdr = ["Tab", "Source System", "What It Holds", "Feeds These Signals", "Rows"]
    for i, h in enumerate(hdr, start=1):
        c = readme.cell(row=7, column=i, value=h)
        c.font, c.fill = HEADER_FONT, HEADER_FILL
    catalog = [
        ("SF Accounts", "Salesforce", "Accounts: segment, ARR, owner, contacts, exec activity", "contact_departed, exec_sponsor_gap, single_threaded; importance"),
        ("SF Contacts", "Salesforce", "People at each customer: role, email, last activity", "contact ranking, email drafts, suggested steps"),
        ("SF Contracts", "Salesforce", "Contracts and their end (renewal) dates", "renewal_approaching"),
        ("SF QBR Events", "Salesforce", "Latest QBR per account", "qbr_overdue, no_contact_90d"),
        ("Payments weekly", "Payments data warehouse", "Weekly volume, auth rate, chargebacks, failed payouts", "volume_drop, volume_growth, auth_rate_drop, chargeback_spike, failed_payout"),
        ("Payment methods", "Payments data warehouse", "Methods enabled per merchant, with dates", "payment_method_gap"),
        ("Support tickets", "Support desk", "Sev-1 tickets", "sev1_open"),
        ("Support weekly", "Support desk", "Open ticket counts per week", "(context)"),
        ("Gong calls", "Gong", "Call briefs, attendees, extracted risks/opportunities, sentiment", "cross_sell_fit, sentiment_decline, no_contact_90d; contact call history"),
        ("Slack messages", "Slack (shared channels)", "Customer messages, classified", "contact_departed"),
        ("NPS responses", "NPS survey", "Relationship NPS by contact", "nps_detractor"),
        ("Sync log", "Ingestion log", "Every successful sync per source", "freshness → 'data incomplete' badges"),
        ("Source freshness", "Computed here", "Newest sync and stale flag per source", "(formulas)"),
        ("Signals (output)", "PriorityOne", "What the tracker detected, with history", "OUTPUT"),
        ("Weekly lists (output)", "PriorityOne", "This week's scored and ranked accounts", "OUTPUT"),
    ]
    for i, (tab, src, what, feeds_) in enumerate(catalog, start=8):
        for j, v in enumerate((tab, src, what, feeds_), start=1):
            readme.cell(row=i, column=j, value=v).font = BODY
        if tab == "Source freshness":
            readme.cell(row=i, column=5, value="–").font = BODY
        else:
            c = readme.cell(row=i, column=5, value=f"=COUNTA('{tab}'!A:A)-1")
            c.font = LINK
        readme.cell(row=i, column=1).hyperlink = f"#'{tab}'!A1"
        readme.cell(row=i, column=1).font = URL
    end = 8 + len(catalog)
    readme[f"A{end + 1}"] = ("Connected tools that don't feed the tracker (Clay enrichment, Looker dashboards) have no tab. "
                             "Gong is deliberately 12 days stale, which is why some accounts show 'data incomplete'.")
    readme[f"A{end + 1}"].font = Font(name=FONT, size=9, italic=True)
    for col, w in zip("ABCDE", (24, 26, 52, 66, 8)):
        readme.column_dimensions[col].width = w

    wb.calculation.fullCalcOnLoad = True  # formulas compute when the file is opened
    wb.save(path)
    s.close()
    print(f"Wrote {path}: {len(wb.sheetnames)} tabs")


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "priorityone_source_data.xlsx")
