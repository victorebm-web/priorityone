"""The AM's 1:1 document as a Word (.docx) file. Takes the dict from main.one_on_one_doc."""
from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

NAVY = RGBColor(0x00, 0x12, 0x22)
GREY = RGBColor(0x64, 0x6F, 0x79)
RED = RGBColor(0x9F, 0x12, 0x39)


def _money(v) -> str:
    v = float(v or 0)
    if v >= 1_000_000:
        return f"${v / 1_000_000:.1f}M"
    if v >= 1_000:
        return f"${v / 1_000:.0f}k"
    return f"${v:,.0f}"


def _shade(cell, hex_fill: str) -> None:
    pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    pr.append(shd)


def _table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold, r.font.size = True, Pt(9)
        _shade(c, "F4F5F6")
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            parts = val if isinstance(val, list) else [val]
            for j, part in enumerate(parts):
                p = cells[i].paragraphs[0] if j == 0 else cells[i].add_paragraph()
                text, style = part if isinstance(part, tuple) else (part, "")
                run = p.add_run(text)
                run.font.size = Pt(9)
                run.bold = "b" in style
                if "grey" in style:
                    run.font.color.rgb = GREY
                if "red" in style:
                    run.font.color.rgb = RED
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Pt(w)
    return t


def _bullets(doc, lines, empty):
    if not lines:
        p = doc.add_paragraph(empty)
        p.runs[0].font.color.rgb = GREY
        return
    for line in lines:
        doc.add_paragraph(line, style="List Bullet")


def one_on_one_docx(d: dict, notes: str = "") -> bytes:
    labels = d["week_status_labels"]
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].font.color.rgb = NAVY

    user, mgr = d["user"], d["manager"]
    week = f"Week of {d['selected']:%d %b %Y}"
    doc.add_heading(f"1:1 preparation: {user.name}" + (f" with {mgr.name}" if mgr else ""), level=1)
    p = doc.add_paragraph(f"{week} · prepared {d['today_d']:%d %b %Y} · from PriorityOne")
    p.runs[0].font.color.rgb = GREY

    overdue = len(d["sm"]["overdue"])
    _table(doc, ["Need attention", "Revenue at risk", "Open Salesforce Tasks", "Left from last week"], [[
        f"{d['need_n']} of {d['book_size']}", _money(d["sm"]["revenue_at_risk"]),
        f"{len(d['open_tasks'])}" + (f" ({overdue} overdue)" if overdue else ""), str(len(d["left_over"]))]])

    doc.add_heading("Where I need help", level=2)
    _bullets(doc, d["asks"], "Nothing waiting on anyone this week.")

    c = d["counts"]
    doc.add_heading("This week's accounts", level=2)
    summary = doc.add_paragraph(" · ".join(f"{c[k]} {labels[k].lower()}" for k in d["status_order"]))
    summary.runs[0].font.color.rgb = GREY
    rows = []
    for pr in d["rows"]:
        a = pr["account"]
        done = [(i["text"] + (f" (overdue since {i['overdue']:%d %b})" if i["overdue"] else ""), "red" if i["overdue"] else "")
                for i in pr["done_items"]] or [("Nothing yet", "red")]
        rows.append([[(a.name, "b"), (f"{a.segment} · {_money(a.arr_usd)} ARR", "grey")], labels[pr["ws"]["status"]],
                     pr["top"]["sentence"] if pr["top"] else "–", done,
                     "–" if pr["ws"]["status"] == "completed" else pr["next"]])
    if rows:
        _table(doc, ["Account", "Status", "Problem", "Done so far", "Next step"], rows, widths=[90, 60, 120, 130, 110])
    else:
        doc.add_paragraph("Nothing on the list this week.")

    lc = d["last_counts"]
    doc.add_heading(f"Last week (week of {d['prev']:%d %b})", level=2)
    if d["last_rows"]:
        summary = doc.add_paragraph(" · ".join(f"{lc[k]} {labels[k].lower()}" for k in d["status_order"]))
        summary.runs[0].font.color.rgb = GREY
        _table(doc, ["Account", "How it ended", "Outcome"],
               [[(pr["account"].name, "b"), labels[pr["ws"]["status"]], pr["ws"]["note"] or pr["next"]] for pr in d["last_rows"]],
               widths=[120, 110, 280])
    else:
        doc.add_paragraph("Nothing was on the list.")

    doc.add_heading("Wins", level=2)
    _bullets(doc, d["wins"], "None logged since last week.")
    if d["renewals"]:
        doc.add_heading("Renewals in the next 60 days", level=2)
        _bullets(doc, [f"{a.renewal_date:%d %b}: {a.name} ({a.segment}, {_money(a.arr_usd)} ARR, in "
                       f"{(a.renewal_date - d['selected']).days} days)" for a in d["renewals"]], "None.")
    if d["wrong_lines"]:
        doc.add_heading("Where I think the ranking was wrong", level=2)
        _bullets(doc, d["wrong_lines"], "")
    doc.add_heading("My notes", level=2)
    for para in (notes.strip() or "–").split("\n"):
        doc.add_paragraph(para)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
